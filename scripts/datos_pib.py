"""
datos_pib.py — Monitor de Actividad Económica: cuadros del PIB trimestral del INE → data/*.json

    python scripts/datos_pib.py            # lee datos/ y escribe data/
    python scripts/datos_pib.py --revisar  # sólo lee y verifica; no escribe

Los gráficos (embed/*.html) leen estos JSON con PM.cargar: ningún número del PIB se escribe a
mano en un gráfico ni en la página. Escribe:

  data/pib_actividad.json   PIB por actividad económica, año de referencia 2017 (SCN 2008)
  data/pib_gasto.json       PIB por tipo de gasto, año de referencia 2017
  data/pib_serie_larga.json crecimiento del PIB con la base 1990 (1991-2024) y la base 2017
  data/kpis.json            las cifras de la cabecera de la página del monitor

Los cuadros se leen por RÓTULO (nunca por fila fija: el INE inserta filas) y las columnas por el
encabezado de año y trimestre. Antes de escribir se verifican las identidades del propio INE
(contribuciones que suman el PIB, estructura que suma 100, variaciones que salen de los niveles);
si alguna no cierra, el script aborta SIN tocar data/ (un dato mal leído no se publica).
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
B17 = BASE / "datos" / "pib_trimestral_base2017"
B90 = BASE / "datos" / "pib_trimestral_base1990"
DATA = BASE / "data"

FUENTE = "Instituto Nacional de Estadística (INE) — Cuentas Nacionales Trimestrales, año de referencia 2017"
URL = "https://www.ine.gob.bo/referencia2017/pib_trimestral.html"

# ── Rótulos del INE → clave ──────────────────────────────────────────────────
# (prefijo normalizado, clave). El orden importa: el primero que calza gana.
ACTIVIDAD = [
    ("PRODUCTO INTERNO BRUTO", "pib"),
    ("DERECHOS", "impuestos"),
    ("VALOR AGREGADO BRUTO", "vab"),
    ("AGRICULTURA", "agro"),
    ("ACTIVIDAD EXTRACTIVA", "extractiva"),
    ("INDUSTRIAS MANUFACTURERAS", "manufactura"),
    ("SUMINISTRO DE ELECTRICIDAD", "electricidad"),
    ("CONSTRUCCION", "construccion"),
    ("COMERCIO", "comercio"),
    ("TRANSPORTE", "transporte"),
    ("ACTIVIDADES DE ALOJAMIENTO", "alojamiento"),
    ("ACTIVIDADES FINANCIERAS", "financieras"),
    ("ADMINISTRACION PUBLICA", "administracion"),
    ("ACTIVIDADES COMUNALES", "comunales"),
]
SECTORES = ["agro", "extractiva", "manufactura", "electricidad", "construccion", "comercio",
            "transporte", "alojamiento", "financieras", "administracion", "comunales"]
NOMBRES = {
    "pib": ("PIB", "Producto interno bruto (a precios de mercado)"),
    "impuestos": ("Impuestos", "Derechos sobre importaciones, IVA no deducible, IT y otros impuestos indirectos"),
    "vab": ("VAB", "Valor agregado bruto"),
    "agro": ("Agropecuaria", "Agricultura, ganadería, silvicultura y pesca"),
    "extractiva": ("Extractiva", "Actividad extractiva (hidrocarburos y minería)"),
    "manufactura": ("Manufactura", "Industrias manufactureras"),
    "electricidad": ("Electricidad y agua", "Suministro de electricidad, gas, agua y recolección de desechos"),
    "construccion": ("Construcción", "Construcción"),
    "comercio": ("Comercio", "Comercio"),
    "transporte": ("Transporte", "Transporte y comunicaciones"),
    "alojamiento": ("Alojamiento", "Actividades de alojamiento y servicio de comidas"),
    "financieras": ("Financieras", "Actividades financieras y de seguros, inmobiliarias y servicios a las empresas"),
    "administracion": ("Adm. pública", "Administración pública, salud y educación de no mercado"),
    "comunales": ("Comunales", "Actividades comunales, sociales y personales"),
}
# Grandes sectores (agregación propia, declarada en el JSON): primario, secundario, terciario
GRANDES = {
    "primario": ("Primario", ["agro", "extractiva"]),
    "secundario": ("Secundario", ["manufactura", "electricidad", "construccion"]),
    "terciario": ("Terciario", ["comercio", "transporte", "alojamiento", "financieras", "administracion", "comunales"]),
}
GASTO = [
    ("PRODUCTO", "pib"),
    ("GASTO DE CONSUMO FINAL DE LA ADMINISTRACION", "gobierno"),
    ("GASTO DE CONSUMO FINAL DE LOS HOGARES", "hogares"),
    ("VARIACION DE EXISTENCIAS", "existencias"),
    ("FORMACION BRUTA DE CAPITAL FIJO", "fbcf"),
    ("FORMACION BRUTA DE CAPITAL", "fbk"),
    ("EXPORTACIONES", "exportaciones"),
    ("MENOS", "importaciones"),
]
NOMBRES_GASTO = {
    "pib": ("PIB", "Producto interno bruto (a precios de mercado)"),
    "gobierno": ("Consumo del gobierno", "Gasto de consumo final de la administración pública"),
    "hogares": ("Consumo de los hogares", "Gasto de consumo final de los hogares e instituciones sin fines de lucro (ISFLSH)"),
    "fbk": ("Formación bruta de capital", "Formación bruta de capital (inversión fija + variación de existencias + objetos valiosos)"),
    "fbcf": ("Inversión fija (FBCF)", "Formación bruta de capital fijo"),
    "existencias": ("Existencias y otros", "Variación de existencias y adquisición menos disposición de objetos valiosos (FBK − FBCF)"),
    "exportaciones": ("Exportaciones", "Exportaciones de bienes y servicios"),
    "importaciones": ("Importaciones", "Importaciones de bienes y servicios"),
}
ROMANO = {"I": 1, "II": 2, "III": 3, "IV": 4}


def norma(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


def num(v):
    if v is None or isinstance(v, str) and not v.strip():
        return None
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return None if x != x else x


class Falla(Exception):
    pass


def leer_cuadro(ruta: Path, mapa) -> tuple[list[str], dict]:
    """Cuadro transpuesto del INE (rubros en filas, trimestres en columnas) → (claves, {clave: valores})."""
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    # fila de trimestres: la que tiene más celdas I/II/III/IV; la de años, la anterior
    def romanos(f):
        return sum(1 for v in f if norma(v) in ROMANO)
    iq = max(range(len(filas)), key=lambda i: romanos(filas[i]))
    if romanos(filas[iq]) < 8:
        raise Falla(f"{ruta.name}: no se encontró la fila de trimestres")
    anios, actual = filas[iq - 1], None
    cols, claves = [], []
    for c, v in enumerate(filas[iq]):
        a = anios[c] if c < len(anios) else None
        m = re.match(r"\s*(\d{4})", str(a)) if a is not None else None
        if m:
            actual = int(m.group(1))
        q = ROMANO.get(norma(v))
        if q and actual:
            cols.append(c)
            claves.append(f"{actual}-T{q}")
    # continuidad: trimestres seguidos, sin saltos ni repetidos
    for a, b in zip(claves, claves[1:]):
        ya, qa = int(a[:4]), int(a[-1])
        if (ya * 4 + qa) + 1 != int(b[:4]) * 4 + int(b[-1]):
            raise Falla(f"{ruta.name}: trimestres no consecutivos ({a} → {b})")
    out = {}
    for f in filas[iq + 1:]:
        rot = next((x for x in f[:3] if isinstance(x, str) and x.strip()), None)
        if not rot:
            continue
        n = norma(rot)
        if n.startswith("FUENTE") or n.startswith("(") or n.startswith("NOTA"):
            break
        clave = next((k for pre, k in mapa if n.startswith(pre) or n.lstrip("- ").startswith(pre)), None)
        if clave is None or clave in out:
            continue
        out[clave] = [num(f[c]) if c < len(f) else None for c in cols]
    return claves, out


def alinear(claves_ref, claves, valores):
    pos = {k: i for i, k in enumerate(claves)}
    return [valores[pos[k]] if k in pos else None for k in claves_ref]


def r(v, d):
    return None if v is None else round(v, d)


def serie(d, claves_ref, claves, k, dec):
    return [r(x, dec) for x in alinear(claves_ref, claves, d[k])]


def cerca(a, b, tol):
    return a is not None and b is not None and abs(a - b) <= tol


def verificar(cond, msg, fallas):
    if not cond:
        fallas.append(msg)


# ── PIB por actividad ────────────────────────────────────────────────────────
def pib_actividad(fallas):
    arch = {
        "nivel": "encadenadas/07_pib_actividad_encadenadas.xlsx",
        "indice": "encadenadas/08_pib_actividad_indices_encadenados.xlsx",
        "var": "encadenadas/09_pib_actividad_variacion_encadenadas.xlsx",
        "var_acum": "encadenadas/10_pib_actividad_variacion_acum_encadenadas.xlsx",
        "contrib": "encadenadas/11_pib_actividad_contribucion_encadenadas.xlsx",
        "contrib_acum": "encadenadas/12_pib_actividad_contribucion_acum_encadenadas.xlsx",
        "corrientes": "corrientes/01_pib_actividad_corrientes.xlsx",
        "estructura": "corrientes/02_pib_actividad_estructura_corrientes.xlsx",
    }
    leidos = {m: leer_cuadro(B17 / a, ACTIVIDAD) for m, a in arch.items()}
    claves = leidos["nivel"][0]
    todas = ["pib", "impuestos", "vab"] + SECTORES
    for m, (ks, d) in leidos.items():
        faltan = [k for k in todas if k not in d]
        if faltan:
            raise Falla(f"{arch[m]}: faltan rubros {faltan}")
        if ks[-1] != claves[-1]:
            fallas.append(f"{arch[m]} termina en {ks[-1]} y el nivel en {claves[-1]}")
    dec = {"nivel": 2, "corrientes": 2, "indice": 3, "var": 4, "var_acum": 4, "contrib": 4, "contrib_acum": 4, "estructura": 4}
    M = {m: {k: serie(d, claves, ks, k, dec[m]) for k in todas} for m, (ks, d) in leidos.items()}

    # Identidades del INE
    n = len(claves)
    for i in range(n):
        q = claves[i]
        # contribuciones (trimestral y acumulada): sectores + impuestos = PIB
        for m in ("contrib", "contrib_acum"):
            p = M[m]["pib"][i]
            if p is None:
                continue
            s = sum(M[m][k][i] or 0 for k in SECTORES) + (M[m]["impuestos"][i] or 0)
            verificar(cerca(s, p, 0.03), f"{m} {q}: sectores+impuestos {s:.3f} ≠ PIB {p:.3f}", fallas)
        # estructura a precios corrientes: suma 100
        if M["estructura"]["pib"][i] is not None:
            s = sum(M["estructura"][k][i] or 0 for k in SECTORES) + (M["estructura"]["impuestos"][i] or 0)
            verificar(cerca(s, 100, 0.05), f"estructura {q}: suma {s:.3f}", fallas)
        # la variación interanual sale del nivel
        if i >= 4:
            for k in todas:
                a, b, v = M["nivel"][k][i - 4], M["nivel"][k][i], M["var"][k][i]
                if a and b is not None and v is not None:
                    verificar(cerca((b / a - 1) * 100, v, 0.02), f"var {k} {q}: {v:.3f} vs nivel {(b / a - 1) * 100:.3f}", fallas)
        # la acumulada sale de los niveles del año
        if i >= 4 and M["var_acum"]["pib"][i] is not None:
            t = int(q[-1])
            ahora = sum(M["nivel"]["pib"][i - j] for j in range(t))
            antes = sum(M["nivel"]["pib"][i - 4 - j] for j in range(t))
            verificar(cerca((ahora / antes - 1) * 100, M["var_acum"]["pib"][i], 0.02),
                      f"var_acum pib {q}: {M['var_acum']['pib'][i]:.3f} vs niveles {(ahora / antes - 1) * 100:.3f}", fallas)
    # grandes sectores: sumas declaradas (contribución y estructura son aditivas; el nivel encadenado NO)
    grandes = {}
    for g, (nom, ks) in GRANDES.items():
        grandes[g] = {"nombre": nom, "sectores": ks}
        for m in ("contrib", "contrib_acum", "estructura", "corrientes"):
            grandes[g][m] = [None if all(M[m][k][i] is None for k in ks) else r(sum(M[m][k][i] or 0 for k in ks), 4)
                             for i in range(n)]
    out = {
        "metadata": {
            "fuente": FUENTE, "url": URL, "ultimo": claves[-1], "primero": claves[0],
            "unidades": {
                "nivel": "millones de bolivianos encadenados (referencia 2017)",
                "indice": "índice de volumen encadenado (2017 = 100)",
                "var": "variación porcentual respecto al mismo trimestre del año anterior",
                "var_acum": "variación porcentual acumulada en el año respecto al mismo período del año anterior",
                "contrib": "contribución a la variación interanual del PIB, en puntos porcentuales",
                "contrib_acum": "contribución a la variación acumulada del PIB, en puntos porcentuales",
                "corrientes": "millones de bolivianos corrientes",
                "estructura": "participación en el PIB a precios corrientes, en porcentaje",
            },
            "notas": [
                "Las medidas de volumen encadenadas no son aditivas: la suma de las actividades no iguala el PIB.",
                "Contribuciones y estructura sí suman el PIB (incluido el rubro de impuestos).",
                "Los grandes sectores (primario, secundario, terciario) son una agregación de POPULI.",
                "Cifras preliminares del INE, sujetas a revisión.",
            ],
        },
        "trimestres": claves,
        "rubros": [{"k": k, "nombre": NOMBRES[k][0], "nombre_largo": NOMBRES[k][1]} for k in todas],
        "sectores": SECTORES,
        "grandes": grandes,
    }
    out.update(M)
    return out


# ── PIB por tipo de gasto ────────────────────────────────────────────────────
def pib_gasto(fallas, claves_act):
    arch = {
        "nivel": "encadenadas/19_pib_gasto_encadenadas.xlsx",
        "var": "encadenadas/20_pib_gasto_crecimiento_encadenadas.xlsx",
        "var_acum": "encadenadas/21_pib_gasto_crecimiento_acum_encadenadas.xlsx",
        "contrib": "encadenadas/22_pib_gasto_contribucion_encadenadas.xlsx",
        "contrib_acum": "encadenadas/23_pib_gasto_contribucion_acum_encadenadas.xlsx",
        "corrientes": "corrientes/05_pib_gasto_corrientes.xlsx",
        "estructura": "corrientes/06_pib_gasto_estructura_corrientes.xlsx",
    }
    leidos = {m: leer_cuadro(B17 / a, GASTO) for m, a in arch.items()}
    claves = leidos["nivel"][0]
    if claves[-1] != claves_act[-1]:
        fallas.append(f"gasto termina en {claves[-1]} y actividad en {claves_act[-1]}")
    base = ["pib", "gobierno", "hogares", "fbk", "fbcf", "exportaciones", "importaciones"]
    dec = {"nivel": 2, "corrientes": 2, "var": 4, "var_acum": 4, "contrib": 4, "contrib_acum": 4, "estructura": 4}
    M = {}
    for m, (ks, d) in leidos.items():
        faltan = [k for k in base if k not in d]
        if faltan:
            raise Falla(f"{arch[m]}: faltan rubros {faltan}")
        M[m] = {k: serie(d, claves, ks, k, dec[m]) for k in base}
        if "existencias" in d:
            M[m]["_existencias_ine"] = serie(d, claves, ks, "existencias", dec[m])
    n = len(claves)
    # CONTRIBUCIONES: el INE publica la de las importaciones con el signo de su variación (sube si las
    # importaciones crecen) y hay que RESTARLA. Acá se guarda la contribución real (negada): así las
    # barras apiladas suman el PIB. En la acumulada (cuadro 05.05) la FBK viene SIN existencias y las
    # existencias en su propia fila: se recompone FBK = FBK(05.05) + existencias para que las dos
    # tablas digan lo mismo.
    for m in ("contrib", "contrib_acum"):
        M[m]["importaciones"] = [None if v is None else r(-v, 4) for v in M[m]["importaciones"]]
        if m == "contrib_acum":
            ex = M[m].pop("_existencias_ine", None)
            if ex is None:
                raise Falla("contribución acumulada sin la fila de variación de existencias")
            M[m]["fbk"] = [None if a is None else r(a + (b or 0), 4) for a, b in zip(M[m]["fbk"], ex)]
    for m in M:
        M[m].pop("_existencias_ine", None)
    # La variación de la FBK que publica el INE (cuadros 05.02 y 05.03) NO sale de su propio nivel encadenado:
    # sigue a la inversión fija (la calcula sin existencias, aunque la nota al pie diga lo contrario).
    # Para no dejar una cifra que contradice al nivel, esas dos tablas no llevan «fbk»: la inversión es «fbcf».
    for m in ("var", "var_acum"):
        M[m].pop("fbk", None)
    # existencias y otros = FBK − FBCF: aditivo en contribuciones, corrientes y estructura (no en el encadenado)
    for m in ("contrib", "contrib_acum", "corrientes", "estructura"):
        M[m]["existencias"] = [None if a is None or b is None else r(a - b, 4) for a, b in zip(M[m]["fbk"], M[m]["fbcf"])]
    for i in range(n):
        q = claves[i]
        for m in ("contrib", "contrib_acum"):
            p = M[m]["pib"][i]
            if p is None:
                continue
            s = sum(M[m][k][i] or 0 for k in ("gobierno", "hogares", "fbk", "exportaciones", "importaciones"))
            verificar(cerca(s, p, 0.03), f"gasto {m} {q}: componentes {s:.3f} ≠ PIB {p:.3f}", fallas)
        for m in ("corrientes",):
            p = M[m]["pib"][i]
            if p is None:
                continue
            s = sum(M[m][k][i] or 0 for k in ("gobierno", "hogares", "fbk", "exportaciones")) - (M[m]["importaciones"][i] or 0)
            verificar(cerca(s, p, max(1.0, abs(p) * 1e-4)), f"gasto corrientes {q}: C+G+I+X−M {s:.1f} ≠ PIB {p:.1f}", fallas)
        if M["estructura"]["pib"][i] is not None:
            s = sum(M["estructura"][k][i] or 0 for k in ("gobierno", "hogares", "fbk", "exportaciones")) - (M["estructura"]["importaciones"][i] or 0)
            verificar(cerca(s, 100, 0.05), f"gasto estructura {q}: suma {s:.3f}", fallas)
        if i >= 4:
            for k in M["var"]:
                a, b, v = M["nivel"][k][i - 4], M["nivel"][k][i], M["var"][k][i]
                if a and b is not None and v is not None:
                    verificar(cerca((b / a - 1) * 100, v, 0.02), f"gasto var {k} {q}: {v:.3f} vs nivel {(b / a - 1) * 100:.3f}", fallas)
    comps = base + ["existencias"]
    out = {
        "metadata": {
            "fuente": FUENTE, "url": URL, "ultimo": claves[-1], "primero": claves[0],
            "unidades": {
                "nivel": "millones de bolivianos encadenados (referencia 2017)",
                "var": "variación porcentual respecto al mismo trimestre del año anterior",
                "var_acum": "variación porcentual acumulada en el año",
                "contrib": "contribución a la variación interanual del PIB, en puntos porcentuales",
                "contrib_acum": "contribución a la variación acumulada del PIB, en puntos porcentuales",
                "corrientes": "millones de bolivianos corrientes",
                "estructura": "participación en el PIB a precios corrientes, en porcentaje",
            },
            "notas": [
                "PIB = consumo de los hogares + consumo del gobierno + formación bruta de capital + exportaciones − importaciones.",
                "En contrib y contrib_acum las importaciones ya llevan su signo real (negativo cuando crecen): los componentes suman el PIB.",
                "En corrientes, estructura y nivel las importaciones van en positivo (es un monto): se restan.",
                "La formación bruta de capital (fbk) incluye la inversión fija (fbcf) y las existencias; 'existencias' = fbk − fbcf (no existe en el nivel encadenado).",
                "var y var_acum no traen fbk: la variación de la FBK que publica el INE no sale de su nivel (sigue a la inversión fija); la inversión se lee en fbcf.",
                "Cifras preliminares del INE, sujetas a revisión.",
            ],
        },
        "trimestres": claves,
        "componentes": [{"k": k, "nombre": NOMBRES_GASTO[k][0], "nombre_largo": NOMBRES_GASTO[k][1]} for k in comps],
    }
    out.update(M)
    return out


# ── Serie larga: base 1990 + base 2017 ───────────────────────────────────────
def serie_larga(fallas, act):
    ruta = B90 / "actividad" / "04_pib_actividad_variacion_similar.xlsx"
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = [list(f) for f in ws.iter_rows(values_only=True)]
    wb.close()
    ih = next(i for i, f in enumerate(filas) if any(norma(v) == "PERIODO" for v in f))
    col = next(c for c, v in enumerate(filas[ih]) if norma(v).startswith("PIB A PRECIOS DE MERCADO"))
    trim, anual, anio = [], [], None
    for f in filas[ih + 1:]:
        a = norma(f[0])
        if not a:
            continue
        if a.startswith("FUENTE"):
            break
        m = re.match(r"^(\d{4})", a)
        if m:
            anio = int(m.group(1))
            anual.append([str(anio), r(num(f[col]), 4)])
            continue
        m = re.match(r"^(IV|III|II|I)\s+TRIMESTRE", a)
        if m and anio:
            trim.append([f"{anio}-T{ROMANO[m.group(1)]}", r(num(f[col]), 4)])
    if not trim or trim[0][0] != "1991-T1":
        raise Falla("serie base 1990: no empieza en 1991-T1")
    # base 2017: la interanual de cada trimestre y la anual = acumulada al IV trimestre
    k17 = act["trimestres"]
    t17 = [[k, v] for k, v in zip(k17, act["var"]["pib"]) if v is not None]
    a17 = [[k[:4], v] for k, v in zip(k17, act["var_acum"]["pib"]) if v is not None and k.endswith("T4")]
    # la anual de la base 1990 tiene que salir de sus trimestres (promedio no: niveles; acá sólo coherencia de signo y rango)
    for k, v in anual:
        verificar(v is not None and -15 < v < 15, f"base 1990 {k}: crecimiento anual fuera de rango ({v})", fallas)
    return {
        "metadata": {
            "fuente": "INE — Cuentas Nacionales Trimestrales: base 1990 (1991-2024) y año de referencia 2017 (2018 en adelante)",
            "url": URL,
            "unidad": "variación porcentual del PIB a precios constantes respecto al mismo período del año anterior",
            "notas": [
                "Son dos mediciones distintas del mismo PIB: no se empalman. Entre 2018 y 2024 se solapan y difieren.",
                "La base 1990 es la serie que el INE dejó de actualizar al adoptar el año de referencia 2017 (SCN 2008).",
                "La anual de la base 2017 es la variación acumulada al IV trimestre.",
            ],
            "ultimo": t17[-1][0] if t17 else None,
        },
        "base1990": {"trimestral": trim, "anual": anual},
        "base2017": {"trimestral": t17, "anual": a17},
    }


# ── Cifras de la cabecera ────────────────────────────────────────────────────
def kpis(act, gasto, larga):
    k = act["trimestres"]
    i = len(k) - 1
    ult = k[i]
    t = int(ult[-1])
    caen = [s for s in SECTORES if act["var"][s][i] is not None and act["var"][s][i] < 0]
    peor = min(SECTORES, key=lambda s: act["var"][s][i])
    mejor = max(SECTORES, key=lambda s: act["var"][s][i])
    anual = larga["base2017"]["anual"]

    # Cifras con 4 decimales: la página redondea una sola vez al mostrar (1,6538 → «1,7»; con 1,65 salía «1,6»).
    # ── Resumen de la página («Qué dicen los datos»): todo sobre la SUMA DE 4 TRIMESTRES (sin estacionalidad)
    niv = act["nivel"]["pib"]
    s4 = lambda a, j: sum(a[j - x] for x in range(4)) if j >= 3 and all(a[j - x] is not None for x in range(4)) else None
    j19 = k.index(f"2019-T{t}") if f"2019-T{t}" in k else None          # la misma ventana de 4 trimestres en 2019
    jmax = max(range(3, len(k)), key=lambda j: s4(niv, j))
    completos = [v for _, v in anual if v is not None]
    seguidos = 0
    for v in reversed(completos):
        if v < 0:
            seguidos += 1
        else:
            break
    # el componente del gasto y la actividad que más empujaron y más frenaron en lo que va del año
    cg = {c: gasto["contrib_acum"][c][i] for c in ("hogares", "gobierno", "fbcf", "existencias", "exportaciones", "importaciones")}
    cs = {s: act["contrib_acum"][s][i] for s in SECTORES}
    ext = lambda d, f: f(d, key=lambda x: d[x])
    nomg = lambda c: NOMBRES_GASTO[c][0]
    # la caída anual anterior a la racha actual (para decir «antes, la última fue en 2020: −12,7 %»)
    previas = [(a, v) for a, v in anual[: len(anual) - seguidos] if v is not None and v < 0]
    resumen = {
        "anios_caida_seguidos": seguidos,
        "caida_previa": {"anio": int(previas[-1][0]), "valor": r(previas[-1][1], 4)} if previas else None,
        "ultimo_anio_completo": int(anual[-1][0]) if anual else None,
        "pib_vs_2019": r((s4(niv, i) / s4(niv, j19) - 1) * 100, 4) if j19 is not None else None,
        "pib_vs_maximo": {"valor": r((s4(niv, i) / s4(niv, jmax) - 1) * 100, 4), "periodo": k[jmax]},
        "gasto_empuja": {"nombre": nomg(ext(cg, max)), "pp": r(cg[ext(cg, max)], 4),
                         "var": r(gasto["var_acum"].get(ext(cg, max), [None] * (i + 1))[i], 4)},
        "gasto_frena": {"nombre": nomg(ext(cg, min)), "pp": r(cg[ext(cg, min)], 4),
                        "var": r(gasto["var_acum"].get(ext(cg, min), [None] * (i + 1))[i], 4)},
        "inversion_fija_acum": {"var": r(gasto["var_acum"]["fbcf"][i], 4), "pp": r(cg["fbcf"], 4)},
        "hogares_acum": {"var": r(gasto["var_acum"]["hogares"][i], 4), "pp": r(cg["hogares"], 4)},
        "sector_empuja": {"nombre": NOMBRES[ext(cs, max)][0], "pp": r(cs[ext(cs, max)], 4), "var": r(act["var_acum"][ext(cs, max)][i], 4)},
        "sector_frena": {"nombre": NOMBRES[ext(cs, min)][0], "pp": r(cs[ext(cs, min)], 4), "var": r(act["var_acum"][ext(cs, min)][i], 4)},
        "construccion_vs_2019": r((s4(act["nivel"]["construccion"], i) / s4(act["nivel"]["construccion"], j19) - 1) * 100, 4) if j19 is not None else None,
    }
    return {
        "ultimo": ult,
        "pib_acumulado": {
            "valor": r(act["var_acum"]["pib"][i], 4), "anio": int(ult[:4]), "trimestres": t,
            "texto": f"Acumulado {'del año' if t == 4 else 'a T' + str(t)} {ult[:4]}",
        },
        "pib_interanual": {"valor": r(act["var"]["pib"][i], 4), "periodo": ult},
        "pib_anual_anterior": {"valor": anual[-1][1] if t != 4 else (anual[-2][1] if len(anual) > 1 else None),
                               "anio": int(anual[-1][0]) if t != 4 else (int(anual[-2][0]) if len(anual) > 1 else None)},
        "sectores_en_caida": {"n": len(caen), "de": len(SECTORES), "periodo": ult,
                              "lista": [NOMBRES[s][0] for s in caen]},
        "sector_peor": {"nombre": NOMBRES[peor][0], "valor": r(act["var"][peor][i], 4)},
        "sector_mejor": {"nombre": NOMBRES[mejor][0], "valor": r(act["var"][mejor][i], 4)},
        "inversion_fija": {"valor": r(gasto["var"]["fbcf"][i], 4), "periodo": ult},
        "resumen": resumen,
        "fuente": "INE",
    }


def escribir(nombre, obj):
    # sin fecha de generación y sin reescribir lo igual: si el INE no publicó nada, el robot no commitea
    txt = json.dumps(obj, ensure_ascii=False, separators=(",", ":"))
    destino = DATA / nombre
    if destino.exists() and destino.read_text(encoding="utf-8") == txt:
        print(f"  = data/{nombre} sin cambios")
        return
    destino.write_text(txt, encoding="utf-8")
    print(f"  ✓ data/{nombre} ({destino.stat().st_size // 1024} KB)")


def main() -> int:
    fallas = []
    try:
        act = pib_actividad(fallas)
        gas = pib_gasto(fallas, act["trimestres"])
        lar = serie_larga(fallas, act)
    except Falla as e:
        print(f"✗ {e}\n  No se escribió nada.")
        return 1
    if fallas:
        print(f"✗ {len(fallas)} identidad(es) no cierran; no se escribió nada:")
        for f in fallas[:25]:
            print("   ·", f)
        return 1
    kp = kpis(act, gas, lar)
    print(f"PIB {act['metadata']['primero']} → {act['metadata']['ultimo']} · identidades verificadas")
    print(f"  acumulado {kp['pib_acumulado']['texto']}: {kp['pib_acumulado']['valor']}% · interanual {kp['pib_interanual']['valor']}%"
          f" · sectores en caída {kp['sectores_en_caida']['n']}/{kp['sectores_en_caida']['de']}: {', '.join(kp['sectores_en_caida']['lista'])}")
    if "--revisar" in sys.argv:
        return 0
    DATA.mkdir(exist_ok=True)
    escribir("pib_actividad.json", act)
    escribir("pib_gasto.json", gas)
    escribir("pib_serie_larga.json", lar)
    escribir("kpis.json", kp)
    return 0


if __name__ == "__main__":
    sys.exit(main())
